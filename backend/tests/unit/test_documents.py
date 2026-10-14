import io
import zipfile

import pytest

from app.documents.chunking import chunk_blocks
from app.documents.parsers import Block, parse
from app.documents.validation import UploadRejected, sanitize_filename, storage_path, validate_upload

MB = 1024 * 1024


def test_extension_mime_and_magic_bytes_are_checked() -> None:
    ok = validate_upload("notes.md", b"# Hello\n\nworld", "text/markdown", MB)
    assert ok.extension == ".md" and len(ok.storage_name) == 35
    with pytest.raises(UploadRejected):
        validate_upload("evil.exe", b"MZ...", None, MB)
    with pytest.raises(UploadRejected):
        validate_upload("fake.pdf", b"not a pdf", "application/pdf", MB)
    with pytest.raises(UploadRejected):
        validate_upload("x.txt", b"hello", "application/pdf", MB)  # declared type mismatch
    with pytest.raises(UploadRejected):
        validate_upload("big.txt", b"a" * (MB + 1), None, MB)
    with pytest.raises(UploadRejected):
        validate_upload("bin.txt", b"abc\x00def", None, MB)
    with pytest.raises(UploadRejected):
        validate_upload("empty.txt", b"", None, MB)


def test_filenames_are_sanitised_and_paths_cannot_escape() -> None:
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("..\\..\\x.pdf") == "x.pdf"
    assert sanitize_filename("re<port>.pdf") == "re_port_.pdf"
    with pytest.raises(UploadRejected):
        sanitize_filename("..")
    uid = "11111111-1111-1111-1111-111111111111"
    p = storage_path("/data", uid, "0123456789abcdef0123456789abcdef.pdf")
    assert str(p).startswith("/data/uploads/" + uid)
    with pytest.raises(UploadRejected):
        storage_path("/data", uid, "../../../etc/passwd")


def test_zip_containers_with_traversal_are_rejected() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../evil.xml", "x")
    with pytest.raises(UploadRejected):
        validate_upload("doc.docx", buf.getvalue(), None, MB)


def test_parsers_preserve_structure() -> None:
    md = parse(b"# Plan\n\nIntro text.\n\n## Risks\n\n- one\n- two", ".md", "plan.md")
    assert md.title == "Plan" and [b.section for b in md.blocks][-1] == "Risks"
    csv = parse(b"a,b\n1,2\n3,4\n", ".csv", "t.csv")
    assert csv.metadata == {"columns": ["a", "b"], "row_count": 2}
    import docx

    d = docx.Document()
    d.add_heading("Title", 0)
    d.add_heading("Section A", 1)
    d.add_paragraph("Body A")
    out = io.BytesIO()
    d.save(out)
    parsed = parse(out.getvalue(), ".docx", "x.docx")
    assert any(b.section == "Section A" and b.text == "Body A" for b in parsed.blocks)
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, "Page one text.")
    pdf.add_page()
    pdf.multi_cell(0, 6, "Page two text.")
    p = parse(bytes(pdf.output()), ".pdf", "x.pdf")
    assert p.page_count == 2 and {b.page for b in p.blocks} == {1, 2}


def test_chunks_never_cross_page_or_section() -> None:
    blocks = [Block("a " * 300, page=1, section="S1"), Block("b " * 300, page=1, section="S2"), Block("c " * 50, page=2, section="S2")]
    chunks = chunk_blocks(blocks, target_chars=500)
    assert all(len({c.page}) == 1 for c in chunks)
    assert [(c.page, c.section) for c in chunks][:2] == [(1, "S1"), (1, "S1")]
    assert chunks[-1].page == 2
