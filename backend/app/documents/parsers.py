"""Format parsers that keep structure: pages (PDF), headings/sections (DOCX/Markdown), sheets (XLSX)."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None
    kind: str = "paragraph"  # paragraph | heading | table | list


@dataclass
class ParsedDocument:
    title: str
    blocks: list[Block]
    page_count: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)


class ParseError(ValueError):
    pass


def parse(content: bytes, ext: str, filename: str) -> ParsedDocument:
    stem = re.sub(r"\.[^.]+$", "", filename)
    try:
        if ext == ".pdf":
            return _pdf(content, stem)
        if ext == ".docx":
            return _docx(content, stem)
        if ext in (".md", ".markdown"):
            return _markdown(_decode(content), stem)
        if ext == ".txt":
            return _plain(_decode(content), stem)
        if ext == ".csv":
            return _csv(_decode(content), stem)
        if ext == ".xlsx":
            return _xlsx(content, stem)
    except ParseError:
        raise
    except Exception as exc:  # malformed files must not crash the pipeline
        raise ParseError(f"could not parse {ext} file: {type(exc).__name__}") from exc
    raise ParseError(f"no parser for {ext}")


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("undecodable text")


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def _pdf(content: bytes, stem: str) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise ParseError("encrypted PDFs are not supported")
    blocks: list[Block] = []
    section: str | None = None
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for para in _paragraphs(text.replace("\r", "")):
            lines = para.splitlines()
            first = lines[0].strip()
            if len(lines) == 1 and len(first) < 80 and (first.isupper() or re.match(r"^(\d+(\.\d+)*\.?\s+)?[A-Z][\w ,&/\-]{2,70}$", first)) and not first.endswith("."):
                section = first
                blocks.append(Block(first, page=i, section=section, kind="heading"))
            else:
                blocks.append(Block(" ".join(ln.strip() for ln in lines), page=i, section=section))
    meta = reader.metadata  # DocumentInformation | None; an empty mapping is falsy just like None
    title = (getattr(meta, "title", None) or stem) if meta else stem
    return ParsedDocument(title=str(title), blocks=blocks, page_count=len(reader.pages),
                          metadata={"author": str(getattr(meta, "author", "") or "") if meta else ""})


def _docx(content: bytes, stem: str) -> ParsedDocument:
    import docx

    d = docx.Document(io.BytesIO(content))
    blocks: list[Block] = []
    section: str | None = None
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            section = text
            blocks.append(Block(text, section=section, kind="heading"))
        elif "list" in style:
            blocks.append(Block("- " + text, section=section, kind="list"))
        else:
            blocks.append(Block(text, section=section))
    for t in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        if rows:
            blocks.append(Block("\n".join(rows), section=section, kind="table"))
    cp = d.core_properties
    return ParsedDocument(title=cp.title or stem, blocks=blocks,
                          metadata={"author": cp.author or "", "created": cp.created.isoformat() if cp.created else None})


def _markdown(text: str, stem: str) -> ParsedDocument:
    blocks: list[Block] = []
    section: str | None = None
    title = stem
    for para in _paragraphs(text):
        h = re.match(r"^(#{1,6})\s+(.+)$", para.splitlines()[0])
        if h:
            section = h.group(2).strip()
            if h.group(1) == "#" and title == stem:
                title = section
            blocks.append(Block(section, section=section, kind="heading"))
            rest = "\n".join(para.splitlines()[1:]).strip()
            if rest:
                blocks.append(Block(rest, section=section))
        else:
            kind = "list" if re.match(r"^\s*([-*+]|\d+\.)\s", para) else "paragraph"
            blocks.append(Block(para, section=section, kind=kind))
    return ParsedDocument(title=title, blocks=blocks)


def _plain(text: str, stem: str) -> ParsedDocument:
    return ParsedDocument(title=stem, blocks=[Block(p) for p in _paragraphs(text)])


def _csv(text: str, stem: str) -> ParsedDocument:
    rows = list(csv.reader(io.StringIO(text)))
    header = rows[0] if rows else []
    preview = rows[1:21]
    lines = [", ".join(header)] + [", ".join(r) for r in preview]
    meta = {"columns": header, "row_count": max(0, len(rows) - 1)}
    blocks = [Block(f"Table {stem} with {meta['row_count']} rows and columns: {', '.join(header)}", kind="table"),
              Block("\n".join(lines), kind="table")]
    return ParsedDocument(title=stem, blocks=blocks, metadata=meta)


def _xlsx(content: bytes, stem: str) -> ParsedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    blocks: list[Block] = []
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 21:
                break
            rows.append(", ".join("" if v is None else str(v) for v in row))
        sheets.append({"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column})
        if rows:
            blocks.append(Block(f"Sheet {ws.title} ({ws.max_row} rows): columns {rows[0]}", section=ws.title, kind="table"))
            blocks.append(Block("\n".join(rows[1:]), section=ws.title, kind="table"))
    wb.close()
    return ParsedDocument(title=stem, blocks=blocks, metadata={"sheets": sheets})
